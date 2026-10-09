"""
Tests for the generated SLL workbook.

The workbook is formula-driven so a reader can change the utilisation and watch
the economics move. That only has value if the formulas agree with the Python
model they were generated from, so :func:`test_excel_formulas_match_python`
evaluates the workbook and compares cell by cell.

That test needs the optional ``formulas`` package. It is skipped rather than
failed when absent, so the suite stays installable with the project's own
requirements; the structural tests above it run either way.
"""

from __future__ import annotations

import os
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

import build_model
import sll

ROOT = Path(__file__).resolve().parent.parent
COMMITTED = ROOT / "project2-sll-structuring" / "SLL_Structuring_Model.xlsx"

RATCHET = "3. Margin Ratchet"
SENSITIVITY = "4. Utilisation Sensitivity"


@pytest.fixture(scope="module")
def built(tmp_path_factory) -> Path:
    """Build the workbook into a temp directory."""
    out = tmp_path_factory.mktemp("sll") / "model.xlsx"
    build_model.build(str(out))
    return out


# ── structure ────────────────────────────────────────────────────────────────

def test_build_produces_a_valid_file(built):
    assert built.stat().st_size > 5_000
    assert zipfile.ZipFile(built).testzip() is None


def test_expected_sheets(built):
    names = load_workbook(built).sheetnames
    assert len(names) == 6
    for expected in ["Borrower Profile", "KPI Tracker", "SPT Calibration",
                     "Margin Ratchet", "Utilisation Sensitivity", "Economics"]:
        assert any(expected in n for n in names), f"missing sheet: {expected}"


def test_merged_ranges_do_not_overlap(built):
    """Excel refuses to open a file with overlapping merged ranges."""
    wb = load_workbook(built)
    for ws in wb.worksheets:
        boxes = []
        for mr in ws.merged_cells.ranges:
            b = range_boundaries(str(mr))
            for other, name in boxes:
                overlap = not (b[2] < other[0] or b[0] > other[2]
                               or b[3] < other[1] or b[1] > other[3])
                assert not overlap, f"{ws.title}: {mr} overlaps {name}"
            boxes.append((b, str(mr)))


def test_model_is_formula_driven_not_a_value_dump(built):
    """
    The workbook must contain live formulas, otherwise the claim that inputs
    recalculate is false.
    """
    wb = load_workbook(built)
    formulas = [c.value for ws in wb.worksheets for row in ws.iter_rows()
                for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) > 40


def test_committed_workbook_is_regenerable(tmp_path):
    """The committed file must come from the committed script."""
    if not COMMITTED.exists():
        pytest.skip("committed workbook not present")
    fresh = tmp_path / "fresh.xlsx"
    build_model.build(str(fresh))
    assert load_workbook(fresh).sheetnames == load_workbook(COMMITTED).sheetnames


def test_no_proprietary_data_references(built):
    """Project 2 must stay free of paid-vendor sourcing, like the rest of the repo."""
    import re
    vendors = re.compile(r"bloomberg|refinitiv|eikon|factset|capital\s*iq", re.I)
    wb = load_workbook(built)
    hits = [f"{ws.title}!{c.coordinate}" for ws in wb.worksheets
            for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and vendors.search(c.value)]
    assert not hits, f"proprietary references: {hits}"


def test_workbook_states_the_borrower_is_fictional(built):
    """A hypothetical borrower must be unmistakably labelled as such."""
    wb = load_workbook(built)
    text = " ".join(c.value for ws in wb.worksheets for row in ws.iter_rows()
                    for c in row if isinstance(c.value, str)).lower()
    assert "fictional" in text
    assert "illustrative" in text


def test_workbook_explains_the_utilisation_caveat(built):
    """
    The point of the new sheets is that the headline saving assumes full
    drawdown. That caveat must be stated in the workbook, not only in the README.
    """
    wb = load_workbook(built)
    text = " ".join(c.value for ws in wb.worksheets for row in ws.iter_rows()
                    for c in row if isinstance(c.value, str)).lower()
    assert "utilisation" in text
    assert "commitment fee" in text
    assert "break-even" in text or "breakeven" in text


# ── the formulas must agree with the model ───────────────────────────────────

def test_excel_formulas_match_python(built):
    """
    Evaluate every headline formula and compare against ``sll.Facility``.

    A formula-driven workbook that disagrees with the code is worse than no
    workbook, because a reader has no reason to doubt it.
    """
    formulas = pytest.importorskip(
        "formulas", reason="optional dependency; install with: pip install formulas")

    model = formulas.ExcelModel().loads(str(built)).finish()
    solution = model.calculate()

    basename = os.path.basename(str(built)).upper()
    values = {}
    for key, cell in solution.items():
        try:
            values[key.upper()] = cell.value[0, 0]
        except Exception:
            continue

    def cell(sheet: str, addr: str):
        return values.get(f"'[{basename}]{sheet.upper()}'!{addr.upper()}")

    f = sll.Facility(utilisation=1.0)

    # Derived pricing block
    for addr, expected in [
        ("D17", f.drawn_m),
        ("D18", f.undrawn_m),
        ("D19", f.commitment_fee_bps),
        ("D20", f.ratchet.adjustment_bps(3)),
        ("D21", f.margin_bps(3)),
        ("D22", f.all_in_rate_pct(3)),
        ("D23", f.drawn_cost_m(3)),
        ("D24", f.undrawn_cost_m()),
        ("D25", f.total_cost_m(3)),
    ]:
        actual = cell(RATCHET, addr)
        assert actual is not None, f"{RATCHET}!{addr} did not evaluate"
        assert float(actual) == pytest.approx(expected, abs=1e-6), f"at {addr}"

    # Scenario table
    for i, met in enumerate([3, 2, 1, 0]):
        row = 29 + i
        assert float(cell(RATCHET, f"F{row}")) == pytest.approx(
            f.total_cost_m(met), abs=1e-6)
        assert float(cell(RATCHET, f"G{row}")) == pytest.approx(
            f.ratchet_benefit_m(met) * 1000, abs=1e-6)

    # Utilisation sensitivity
    for i, expected in enumerate(f.utilisation_sensitivity()):
        row = 6 + i
        assert float(cell(SENSITIVITY, f"C{row}")) == pytest.approx(
            expected["total_cost_base_m"], abs=1e-6)
        assert float(cell(SENSITIVITY, f"D{row}")) == pytest.approx(
            expected["max_saving_m"] * 1000, abs=1e-6)
        assert float(cell(SENSITIVITY, f"F{row}")) == pytest.approx(
            expected["net_benefit_m"] * 1000, abs=1e-6)

    # Break-even
    assert float(cell(SENSITIVITY, "B14")) == pytest.approx(
        f.breakeven_utilisation(), abs=1e-9)

    # KPI progress
    for i, (base, current, target) in enumerate(
            [(310, 267, 217), (18, 34, 60), (35, 52, 80)]):
        assert float(cell("1. KPI Tracker", f"F{4 + i}")) == pytest.approx(
            sll.kpi_progress(base, current, target), abs=1e-9)

    # Glide path
    for i, (_, value) in enumerate(sll.carbon_glide_path(310, 217, 2024, 2028)):
        col = chr(ord("A") + i)
        assert float(cell("2. SPT Calibration", f"{col}16")) == pytest.approx(
            value, abs=1e-9)

    # Glide paths for the other two KPIs, and the compounded-rate comparison
    for row, args in [(27, (18, 60, 2024, 2028)), (32, (35, 80, 2024, 2027))]:
        for i, (_, value) in enumerate(sll.linear_glide_path(*args)):
            col = chr(ord("A") + i)
            assert float(cell("2. SPT Calibration", f"{col}{row}")) == pytest.approx(
                value, abs=1e-9), f"glide path row {row}, column {col}"
    assert float(cell("2. SPT Calibration", "B19")) == pytest.approx(
        sll.compound_annual_rate(310, 217, 4), abs=1e-9)
