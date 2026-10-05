"""
Tests for the generated Excel model.

The portfolio's headline claim is that no proprietary data is used anywhere.
That claim was false for a long time because the committed workbook still
carried terminal-sourcing instructions inside it, where no text search of the
repository would find them. :func:`test_no_proprietary_data_references` is the
regression guard for that.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

import process_data

ROOT = Path(__file__).resolve().parent.parent
PROJECT1 = ROOT / "project1-green-bond-analysis"
COMMITTED_WORKBOOK = PROJECT1 / "Green_Bond_Market_Analysis.xlsx"

# Data vendors whose use the portfolio explicitly disclaims.
PROPRIETARY = re.compile(
    r"bloomberg|refinitiv|eikon|factset|capital\s*iq|\bmarkit\b|datastream", re.I)


AGG_DIR = PROJECT1 / "data" / "aggregates"


@pytest.fixture(scope="module")
def built_workbook(tmp_path_factory):
    """
    Build the model into a temp directory and return the path.

    Built from the committed aggregates, which is also how CI and anyone
    cloning the repository build it — so these tests exercise the real default
    path rather than one that needs data the repository cannot carry.
    """
    out = tmp_path_factory.mktemp("wb") / "model.xlsx"
    process_data.build(str(out), agg_dir=str(AGG_DIR))
    return out


def _all_strings(path: Path):
    """Yield every string cell value in the workbook, with its location."""
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    yield ws.title, cell.coordinate, cell.value


# ── build ────────────────────────────────────────────────────────────────────

def test_build_produces_a_file(built_workbook):
    assert built_workbook.exists()
    assert built_workbook.stat().st_size > 10_000


def test_build_returns_metrics(tmp_path):
    out = tmp_path / "m.xlsx"
    m = process_data.build(str(out), agg_dir=str(AGG_DIR))
    assert m["total_deals"] > 0
    assert m["total_volume_usd_bn"] > 0


def test_workbook_is_a_valid_zip(built_workbook):
    assert zipfile.ZipFile(built_workbook).testzip() is None


def test_expected_sheets_present(built_workbook):
    names = load_workbook(built_workbook).sheetnames
    assert len(names) == 8
    for expected in ["Summary Dashboard", "Data Quality", "Annual Issuance",
                     "Geography", "Sector Breakdown", "Theme Evolution",
                     "SLB Analysis", "Greenium"]:
        assert any(expected in n for n in names), f"missing sheet: {expected}"


@pytest.mark.parametrize("target", ["built", "committed"])
def test_workbook_carries_no_record_level_source_data(built_workbook, target):
    """
    The workbook must not redistribute the source extract.

    It used to: a "Cleaned Data" sheet carried all 732 CBI records, so deleting
    the CSV would have achieved nothing while the workbook still shipped them.
    CBI's terms of use do not permit that, and the sheet is gone — this is the
    guard against it coming back under any name. A row cap rather than a name
    check, so a renamed sheet cannot slip through.
    """
    path = built_workbook if target == "built" else COMMITTED_WORKBOOK
    if not path.exists():
        pytest.skip(f"{path} not present")
    wb = load_workbook(path)
    oversized = [(ws.title, ws.max_row) for ws in wb.worksheets if ws.max_row > 60]
    assert not oversized, (
        "sheet(s) large enough to hold record-level data: "
        + ", ".join(f"{t} ({n} rows)" for t, n in oversized))


def test_merged_ranges_do_not_overlap(built_workbook):
    """Excel refuses to open a file with overlapping merged ranges."""
    wb = load_workbook(built_workbook)
    for ws in wb.worksheets:
        boxes = []
        for mr in ws.merged_cells.ranges:
            b = range_boundaries(str(mr))
            for other, name in boxes:
                overlap = not (b[2] < other[0] or b[0] > other[2]
                               or b[3] < other[1] or b[1] > other[3])
                assert not overlap, f"{ws.title}: {mr} overlaps {name}"
            boxes.append((b, str(mr)))


def test_charts_are_present(built_workbook):
    wb = load_workbook(built_workbook)
    total = sum(len(ws._charts) for ws in wb.worksheets)
    assert total >= 4


# ── the regression this suite exists for ─────────────────────────────────────

@pytest.mark.parametrize("target", ["built", "committed"])
def test_no_proprietary_data_references(built_workbook, target):
    """
    No cell may cite a paid data vendor as a source or an instruction.

    Sentences that say such data is *not* used are allowed; the check looks for
    the vendor names themselves.
    """
    path = built_workbook if target == "built" else COMMITTED_WORKBOOK
    if not path.exists():
        pytest.skip(f"{path} not present")
    hits = [(sheet, coord, value) for sheet, coord, value in _all_strings(path)
            if PROPRIETARY.search(value)]
    assert not hits, "proprietary data references found:\n" + "\n".join(
        f"  {s}!{c}: {v[:120]}" for s, c, v in hits)


def test_greenium_sheet_states_its_sign_convention(built_workbook):
    """
    A greenium figure is meaningless without a stated sign convention, and the
    two artifacts previously used opposite ones.
    """
    wb = load_workbook(built_workbook)
    sheet = next(ws for ws in wb.worksheets if "Greenium" in ws.title)
    text = " ".join(c.value for row in sheet.iter_rows() for c in row
                    if isinstance(c.value, str)).lower()
    assert "sign convention" in text
    assert "negative" in text


def test_greenium_sheet_cites_open_sources(built_workbook):
    """The framework must be executable without a paid subscription."""
    wb = load_workbook(built_workbook)
    sheet = next(ws for ws in wb.worksheets if "Greenium" in ws.title)
    text = " ".join(c.value for row in sheet.iter_rows() for c in row
                    if isinstance(c.value, str)).lower()
    for source in ["finanzagentur", "aft.gouv.fr", "dmo.gov.uk"]:
        assert source in text, f"expected open source {source!r} on the greenium sheet"


def test_committed_workbook_matches_the_current_script(tmp_path):
    """
    The committed workbook must be regenerable from the committed data.

    It previously was not: the file in the repo had entirely different sheets
    from the ones the script produced, so the README described output nobody
    could reproduce. Cell-level rather than sheet-name-level, so a stale
    committed workbook cannot pass on structure alone.
    """
    if not COMMITTED_WORKBOOK.exists():
        pytest.skip("committed workbook not present")
    fresh = tmp_path / "fresh.xlsx"
    process_data.build(str(fresh), agg_dir=str(AGG_DIR))
    a, b = load_workbook(fresh), load_workbook(COMMITTED_WORKBOOK)
    assert a.sheetnames == b.sheetnames
    for name in a.sheetnames:
        rows_a = list(a[name].iter_rows(values_only=True))
        rows_b = list(b[name].iter_rows(values_only=True))
        assert rows_a == rows_b, (
            f"committed workbook is stale on sheet {name!r} — "
            f"rerun scripts/process_data.py")


def test_build_from_raw_matches_build_from_aggregates(raw_extract, tmp_path):
    """
    Both build paths must produce the same workbook.

    This is what makes the committed aggregates trustworthy as a substitute for
    the extract: given the extract, recomputing the tables changes nothing.
    Only runs for whoever holds a licensed copy.
    """
    if raw_extract is None:
        pytest.skip("no licensed copy of the source extract present (expected in CI)")
    from_agg, from_raw = tmp_path / "agg.xlsx", tmp_path / "raw.xlsx"
    process_data.build(str(from_agg), agg_dir=str(AGG_DIR))
    process_data.build(str(from_raw), input_path=str(raw_extract))
    a, b = load_workbook(from_agg), load_workbook(from_raw)
    assert a.sheetnames == b.sheetnames
    for name in a.sheetnames:
        assert (list(a[name].iter_rows(values_only=True))
                == list(b[name].iter_rows(values_only=True))), f"sheet {name} differs"


def test_dashboard_reports_the_dataset_scope(built_workbook):
    """
    The dashboard must state the real year range and reconcile its total, since
    an earlier version headlined 2018-2024 while totalling 2015-2024 volume.
    """
    wb = load_workbook(built_workbook)
    ws = wb["0. Summary Dashboard"]
    text = " ".join(c.value for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str))
    assert "2015" in text and "2024" in text
    assert "reconcile" in text.lower()
    assert "partial year" in text.lower()
